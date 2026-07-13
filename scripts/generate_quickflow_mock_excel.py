from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime, timedelta
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "src" / "modules" / "metrics" / "mock-throughput-data.ts"
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT_FILE = OUTPUT_DIR / "quickflow-jira-mock.xlsx"
TODAY = datetime(2026, 7, 12, 10, 0, 0)


ISSUE_PATTERN = re.compile(
    r"\{ key: '([^']+)', summary: '([^']+)', issueType: '([^']+)', status: '([^']+)', "
    r"resolvedAt: '([^']+)', assignee: '([^']+)', squad: '([^']+)', priority: '([^']+)', "
    r"labels: \[([^\]]*)\] \}"
)


def parse_mock_issues() -> list[dict]:
  content = SOURCE_FILE.read_text(encoding="utf-8")
  issues = []

  for match in ISSUE_PATTERN.finditer(content):
    labels = [label.strip().strip("'") for label in match.group(9).split(",") if label.strip()]
    issues.append(
        {
            "key": match.group(1),
            "summary": match.group(2),
            "issue_type": match.group(3),
            "status": match.group(4),
            "resolved_at": datetime.fromisoformat(match.group(5).replace("Z", "+00:00")).replace(tzinfo=None),
            "assignee": match.group(6),
            "squad": match.group(7),
            "priority": match.group(8),
            "labels": labels,
        }
    )

  return issues


def build_done_rows(issues: list[dict]) -> list[dict]:
  rows = []
  for index, issue in enumerate(issues):
    lead_days = 4 + (index % 6)
    backlog_days = 2 + (index % 5)
    dev_days = 1 + (index % 4)
    test_wait_days = 1
    test_days = 1 + (index % 3)
    deploy_days = 0

    done_at = issue["resolved_at"]
    deploy_to_prod_at = done_at - timedelta(hours=6)
    in_test_at = deploy_to_prod_at - timedelta(days=test_days)
    ready_for_test_at = in_test_at - timedelta(days=test_wait_days)
    in_dev_at = ready_for_test_at - timedelta(days=dev_days)
    ready_for_dev_at = done_at - timedelta(days=lead_days)
    if in_dev_at < ready_for_dev_at:
      in_dev_at = ready_for_dev_at + timedelta(hours=8)
    created_at = ready_for_dev_at - timedelta(days=backlog_days)
    updated_at = done_at

    rows.append(
        {
            "Issue Key": issue["key"],
            "Summary": issue["summary"],
            "Issue Type": issue["issue_type"],
            "Status": "Done",
            "Priority": issue["priority"],
            "Squad": issue["squad"],
            "Assignee": issue["assignee"],
            "Labels": ", ".join(issue["labels"]),
            "Created": created_at,
            "Updated": updated_at,
            "Resolved": done_at,
            "Ready for Development": ready_for_dev_at,
            "In Development": in_dev_at,
            "Ready for Test": ready_for_test_at,
            "In Test": in_test_at,
            "Deploy to Production": deploy_to_prod_at,
            "Done": done_at,
            "Current Column": "Em Produção",
            "Current Column Entered At": done_at,
            "Lead Time (days)": round((done_at - ready_for_dev_at).total_seconds() / 86400, 2),
            "Cycle Time (days)": round((done_at - in_dev_at).total_seconds() / 86400, 2),
            "Aging WIP (days)": 0,
            "Blocked Days": index % 3,
            "Story Points": 2 + (index % 8),
            "Source": "Jira Mock",
            "Board": "KAN",
        }
    )

  return rows


def build_open_rows(start_index: int) -> list[dict]:
  specs = [
      ("KAN-201", "Consolidar relatorio executivo de julho", "Story", "High", "Core", "Joao", "Em Desenvolvimento", 7),
      ("KAN-202", "Ajustar importacao de CSV do Jira", "Task", "Medium", "Platform", "Ana", "Pronto para Desenvolvimento", 3),
      ("KAN-203", "Investigar erro intermitente no parser", "Bug", "High", "Core", "Carlos", "Em Testes", 5),
      ("KAN-204", "Criar visao de aging por squad", "Story", "Medium", "Growth", "Marina", "Pronto para Testes", 4),
      ("KAN-205", "Preparar filtros por prioridade", "Task", "Low", "Growth", "Marina", "Em Desenvolvimento", 10),
      ("KAN-206", "Automatizar carga de planilhas", "Spike", "Medium", "Platform", "Ana", "Em Deploy para Produção", 2),
      ("KAN-207", "Refinar legenda dos percentis", "Task", "Low", "Core", "Joao", "Pronto para Desenvolvimento", 6),
      ("KAN-208", "Validar dashboard com operacoes", "Story", "High", "Growth", "Carlos", "Em Testes", 8),
      ("KAN-209", "Padronizar colunas de exportacao", "Task", "Medium", "Platform", "Ana", "Em Desenvolvimento", 12),
      ("KAN-210", "Montar amostra para treinamento interno", "Story", "Low", "Core", "Marina", "Pronto para Testes", 9),
  ]

  rows = []
  for offset, spec in enumerate(specs):
    key, summary, issue_type, priority, squad, assignee, current_column, aging_days = spec
    created_at = TODAY - timedelta(days=aging_days + 5 + (offset % 4))
    ready_for_dev_at = TODAY - timedelta(days=aging_days + max(1, offset % 3))
    in_dev_at = ready_for_dev_at + timedelta(days=1)
    ready_for_test_at = in_dev_at + timedelta(days=2 + (offset % 2))
    in_test_at = ready_for_test_at + timedelta(days=1)
    deploy_at = in_test_at + timedelta(days=1)

    current_entered_at_map = {
        "Pronto para Desenvolvimento": ready_for_dev_at,
        "Em Desenvolvimento": in_dev_at,
        "Pronto para Testes": ready_for_test_at,
        "Em Testes": in_test_at,
        "Em Deploy para Produção": deploy_at,
    }

    current_entered_at = current_entered_at_map[current_column]
    rows.append(
        {
            "Issue Key": key,
            "Summary": summary,
            "Issue Type": issue_type,
            "Status": current_column,
            "Priority": priority,
            "Squad": squad,
            "Assignee": assignee,
            "Labels": "aging, wip",
            "Created": created_at,
            "Updated": TODAY - timedelta(days=max(0, aging_days - 1)),
            "Resolved": None,
            "Ready for Development": ready_for_dev_at,
            "In Development": in_dev_at if current_column != "Pronto para Desenvolvimento" else None,
            "Ready for Test": ready_for_test_at if current_column in {"Pronto para Testes", "Em Testes", "Em Deploy para Produção"} else None,
            "In Test": in_test_at if current_column in {"Em Testes", "Em Deploy para Produção"} else None,
            "Deploy to Production": deploy_at if current_column == "Em Deploy para Produção" else None,
            "Done": None,
            "Current Column": current_column,
            "Current Column Entered At": current_entered_at,
            "Lead Time (days)": round((TODAY - ready_for_dev_at).total_seconds() / 86400, 2),
            "Cycle Time (days)": round((TODAY - in_dev_at).total_seconds() / 86400, 2) if in_dev_at else None,
            "Aging WIP (days)": round((TODAY - current_entered_at).total_seconds() / 86400, 2),
            "Blocked Days": 1 + (offset % 4),
            "Story Points": 3 + (offset % 5),
            "Source": "Jira Mock",
            "Board": "KAN",
        }
    )

  return rows


def add_readme_sheet(workbook: Workbook) -> None:
  ws = workbook.active
  ws.title = "README"
  rows = [
      ["QuickFlow Mock Dataset"],
      ["Arquivo criado para testar importacao do QuickFlow com dados semelhantes a um export do Jira."],
      [""],
      ["Planilhas"],
      ["README", "Guia rapido de uso."],
      ["Jira_Mock_Data", "Base principal para importacao, com itens concluidos e itens em WIP."],
      ["Expected_Reference", "Metricas de referencia para comparar o resultado do site."],
      [""],
      ["Colunas principais"],
      ["Ready for Development", "Marco de compromisso usado para lead time."],
      ["Done", "Data final de entrega em producao."],
      ["Current Column", "Coluna atual do fluxo para teste de aging WIP."],
      ["Aging WIP (days)", "Envelhecimento atual do item em andamento."],
      [""],
      ["Sugestao"],
      ["Importe a aba Jira_Mock_Data como .xlsx no site e confira throughput, cycle time, aging e percentis."],
  ]

  for row in rows:
    ws.append(row)

  ws["A1"].font = Font(size=15, bold=True)
  for cell in ws[4]:
    cell.font = Font(bold=True)
  for cell in ws[9]:
    cell.font = Font(bold=True)
  for cell in ws[15]:
    cell.font = Font(bold=True)
  ws.column_dimensions["A"].width = 28
  ws.column_dimensions["B"].width = 92
  ws.freeze_panes = "A1"


def add_data_sheet(workbook: Workbook, rows: list[dict]) -> None:
  ws = workbook.create_sheet("Jira_Mock_Data")
  headers = list(rows[0].keys())
  ws.append(headers)

  for row in rows:
    ws.append([row.get(header) for header in headers])

  header_fill = PatternFill("solid", fgColor="1D3557")
  header_font = Font(color="FFFFFF", bold=True)

  for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

  date_columns = {
      "Created",
      "Updated",
      "Resolved",
      "Ready for Development",
      "In Development",
      "Ready for Test",
      "In Test",
      "Deploy to Production",
      "Done",
      "Current Column Entered At",
  }

  for col_idx, header in enumerate(headers, start=1):
    max_length = len(header)
    for row_idx in range(2, ws.max_row + 1):
      value = ws.cell(row=row_idx, column=col_idx).value
      if header in date_columns and value:
        ws.cell(row=row_idx, column=col_idx).number_format = "yyyy-mm-dd hh:mm"
      if value is not None:
        max_length = max(max_length, len(str(value)))
    ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 30)

  ws.freeze_panes = "A2"
  table = Table(displayName="JiraMockData", ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
  table.tableStyleInfo = TableStyleInfo(
      name="TableStyleMedium2",
      showFirstColumn=False,
      showLastColumn=False,
      showRowStripes=True,
      showColumnStripes=False,
  )
  ws.add_table(table)


def percentile(values: list[float], ratio: float) -> float:
  sorted_values = sorted(values)
  if not sorted_values:
    return 0
  index = max(0, min(len(sorted_values) - 1, math.ceil(len(sorted_values) * ratio) - 1))
  return sorted_values[index]


def add_reference_sheet(workbook: Workbook, rows: list[dict]) -> None:
  ws = workbook.create_sheet("Expected_Reference")
  ws.append(["Reference", "Value"])

  done_rows = [row for row in rows if row["Done"]]
  done_lead_times = [float(row["Lead Time (days)"]) for row in done_rows]

  month_groups: dict[str, dict[str, list[float] | int]] = defaultdict(lambda: {"throughput": 0, "lead_times": []})
  for row in done_rows:
    month = row["Done"].strftime("%Y-%m")
    month_groups[month]["throughput"] += 1
    month_groups[month]["lead_times"].append(float(row["Lead Time (days)"]))

  ws.append(["Overall P50 (days)", percentile(done_lead_times, 0.50)])
  ws.append(["Overall P85 (days)", percentile(done_lead_times, 0.85)])
  ws.append(["Overall P90 (days)", percentile(done_lead_times, 0.90)])
  ws.append(["Done issues", len(done_rows)])
  ws.append(["Open issues", len(rows) - len(done_rows)])
  ws.append([])
  ws.append(["Month", "Throughput", "Average Lead Time"])

  for month in sorted(month_groups):
    lead_times = month_groups[month]["lead_times"]
    average = round(sum(lead_times) / len(lead_times), 2) if lead_times else 0
    ws.append([month, month_groups[month]["throughput"], average])

  ws["A1"].font = Font(bold=True)
  ws["B1"].font = Font(bold=True)
  ws["A7"].font = Font(bold=True)
  ws["B7"].font = Font(bold=True)
  ws["C7"].font = Font(bold=True)
  ws.column_dimensions["A"].width = 24
  ws.column_dimensions["B"].width = 18
  ws.column_dimensions["C"].width = 22


def main() -> None:
  OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
  base_issues = parse_mock_issues()
  rows = build_done_rows(base_issues) + build_open_rows(len(base_issues))

  workbook = Workbook()
  add_readme_sheet(workbook)
  add_data_sheet(workbook, rows)
  add_reference_sheet(workbook, rows)
  workbook.save(OUTPUT_FILE)
  print(f"Workbook created: {OUTPUT_FILE}")


if __name__ == "__main__":
  main()
